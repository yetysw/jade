"""Genere un corpus de test francais dans testdata/ : docx, xlsx, pdf, txt.

Simule des documents internes d'entreprise (procedures, inventaire, contrat,
notes) avec des faits precis et verifiables pour tester le retrieval.
Usage : python scripts/make_testdocs.py [dossier_cible]
"""

from __future__ import annotations

import sys
from pathlib import Path

TARGET = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent / "testdata"


def make_docx_procedure(path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Procédure de sauvegarde des serveurs", level=1)
    doc.add_paragraph(
        "Cette procédure décrit la sauvegarde quotidienne des serveurs de production "
        "de Fiduciaire Exemple SA. Elle est appliquée par l'équipe informatique."
    )
    doc.add_heading("Planification", level=2)
    doc.add_paragraph(
        "Les sauvegardes complètes sont exécutées chaque nuit à 22h30 via Veeam Backup "
        "& Replication. Les sauvegardes incrémentales tournent toutes les 4 heures en journée."
    )
    doc.add_paragraph(
        "La rétention est de 30 jours sur le dépôt local (NAS Synology RS1221+) et de "
        "90 jours sur la copie hors site chiffrée."
    )
    doc.add_heading("Restauration", level=2)
    doc.add_paragraph(
        "Toute demande de restauration doit être adressée au support via un ticket. "
        "Le délai de restauration garanti (RTO) est de 4 heures pour un fichier isolé "
        "et de 24 heures pour un serveur complet."
    )
    doc.add_heading("Contacts", level=2)
    table = doc.add_table(rows=3, cols=2)
    rows = [
        ("Responsable sauvegardes", "Marc Dupont, interne 245"),
        ("Astreinte soir/week-end", "079 555 12 34"),
        ("Escalade fournisseur", "support@example.com"),
    ]
    for row, (k, v) in zip(table.rows, rows):
        row.cells[0].text = k
        row.cells[1].text = v
    doc.save(str(path))


def make_docx_conges(path: Path) -> None:
    import docx

    doc = docx.Document()
    doc.add_heading("Règlement interne — Congés et absences", level=1)
    doc.add_heading("Droit aux vacances", level=2)
    doc.add_paragraph(
        "Chaque collaborateur à plein temps a droit à 25 jours de vacances par année "
        "civile. Les collaborateurs de plus de 50 ans bénéficient de 27 jours."
    )
    doc.add_heading("Demande de congé", level=2)
    doc.add_paragraph(
        "Les demandes de congé se font dans l'outil Tipee au minimum 2 semaines à "
        "l'avance. Elles sont validées par le supérieur direct. Au-delà de 5 jours "
        "consécutifs, la validation de la direction est également requise."
    )
    doc.add_heading("Absences maladie", level=2)
    doc.add_paragraph(
        "Un certificat médical est exigé dès le 3e jour d'absence. Il doit être remis "
        "aux ressources humaines dans les 5 jours."
    )
    doc.save(str(path))


def make_xlsx_inventaire(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Serveurs"
    ws.append(["Nom", "Rôle", "OS", "IP", "Emplacement", "Garantie jusqu'au"])
    for row in [
        ("SRV-DC01", "Contrôleur de domaine", "Windows Server 2022", "192.168.10.5", "Local technique Genève", "2027-03-15"),
        ("SRV-RDS01", "Serveur RDS bureautique", "Windows Server 2022", "192.168.10.21", "Local technique Genève", "2027-03-15"),
        ("SRV-FILES", "Serveur de fichiers", "Windows Server 2019", "192.168.10.30", "Local technique Genève", "2026-11-01"),
        ("PVE002", "Hyperviseur Proxmox", "Proxmox VE 8", "192.168.20.10", "Datacenter Jade", "2028-06-30"),
    ]:
        ws.append(row)
    ws2 = wb.create_sheet("Licences")
    ws2.append(["Produit", "Quantité", "Renouvellement", "Coût annuel CHF"])
    for row in [
        ("Microsoft 365 Business Premium", 42, "2026-09-01", 9450),
        ("Veeam Backup Essentials", 2, "2027-01-15", 1280),
        ("Antivirus Bitdefender GravityZone", 55, "2026-08-20", 2035),
    ]:
        ws2.append(row)
    wb.save(str(path))


def make_pdf_contrat(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm)
    story = [
        Paragraph("Contrat de maintenance informatique", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            "Entre Fiduciaire Exemple SA (le Client) et Jade Sàrl (le Prestataire), "
            "il est convenu ce qui suit.", styles["Normal"]),
        Spacer(1, 8),
        Paragraph("Article 1 — Périmètre", styles["Heading2"]),
        Paragraph(
            "Le Prestataire assure la maintenance préventive et corrective de "
            "l'infrastructure serveur, du parc de postes de travail (55 postes) et "
            "du réseau du Client.", styles["Normal"]),
        Paragraph("Article 2 — Niveaux de service", styles["Heading2"]),
        Paragraph(
            "Temps de prise en charge garanti : 2 heures ouvrées pour un incident "
            "critique (production arrêtée), 8 heures ouvrées pour un incident majeur, "
            "2 jours ouvrés pour une demande standard.", styles["Normal"]),
        Paragraph("Article 3 — Tarification", styles["Heading2"]),
        Paragraph(
            "Le forfait mensuel est de CHF 3'900.- HT, incluant 20 heures "
            "d'interventions. Les heures supplémentaires sont facturées CHF 165.-/h.",
            styles["Normal"]),
        Paragraph("Article 4 — Durée et résiliation", styles["Heading2"]),
        Paragraph(
            "Le contrat est conclu pour 24 mois à compter du 1er janvier 2026, "
            "renouvelable tacitement. Résiliation possible avec un préavis de 3 mois "
            "avant l'échéance.", styles["Normal"]),
    ]
    doc.build(story)


def make_txt_notes(path: Path) -> None:
    path.write_text(
        "Notes de la séance informatique du 12 juin 2026\n"
        "\n"
        "Participants : direction, Jade (M. Kanaan), responsable administratif.\n"
        "\n"
        "Décisions prises :\n"
        "- La migration du serveur de fichiers SRV-FILES vers le cloud est reportée "
        "au premier trimestre 2027, en attendant la fin de la garantie matérielle.\n"
        "- Le Wi-Fi invité sera séparé du réseau interne via un VLAN dédié (VLAN 60), "
        "mise en place prévue en septembre 2026.\n"
        "- Budget validé pour le renouvellement de 12 postes de travail : CHF 18'000.-\n"
        "\n"
        "Prochaine séance : 18 septembre 2026 à 14h00.\n",
        encoding="utf-8",
    )


def make_txt_wifi(path: Path) -> None:
    path.write_text(
        "Procédure : connexion au Wi-Fi d'entreprise\n"
        "\n"
        "Réseau interne : SSID « EXEMPLE-CORP ». L'authentification se fait "
        "automatiquement par certificat machine (802.1X) pour les postes du domaine. "
        "Aucun mot de passe à saisir.\n"
        "\n"
        "Réseau invité : SSID « EXEMPLE-GUEST ». Le mot de passe est affiché à la "
        "réception et change chaque premier lundi du mois. Débit limité à 20 Mbit/s.\n"
        "\n"
        "En cas de problème de connexion, contacter le support Jade au 022 555 44 33 "
        "ou support@example.com.\n",
        encoding="utf-8",
    )


def main() -> None:
    TARGET.mkdir(parents=True, exist_ok=True)
    make_docx_procedure(TARGET / "procedure_sauvegarde.docx")
    make_docx_conges(TARGET / "reglement_conges.docx")
    make_xlsx_inventaire(TARGET / "inventaire_it.xlsx")
    make_pdf_contrat(TARGET / "contrat_maintenance.pdf")
    make_txt_notes(TARGET / "notes_seance_2026-06-12.txt")
    make_txt_wifi(TARGET / "procedure_wifi.txt")
    print(f"Corpus de test genere dans {TARGET} ({len(list(TARGET.iterdir()))} fichiers)")


if __name__ == "__main__":
    main()
